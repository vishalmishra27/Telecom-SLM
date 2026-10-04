from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from fastapi import UploadFile
from openpyxl import load_workbook

from .models import IngestionFileSummary, IngestionResponse


NODE_HEADERS = [
    "canonical_id:ID",
    "name",
    "ontology_class",
    ":LABEL",
    "source_system",
    "source_record_id",
    "source_type",
    "scenario_id",
    "properties_json",
]

REL_HEADERS = [
    ":START_ID",
    ":END_ID",
    ":TYPE",
    "ontology_property",
    "source_system",
    "source_record_id",
    "scenario_id",
    "properties_json",
]

ID_COLUMNS = ("canonical_id", "canonical id", "id", "entity_id", "entity id")
NAME_COLUMNS = ("name", "label", "title", "display_name", "display name")
CLASS_COLUMNS = ("ontology_class", "ontology class", "class", "entity_type", "entity type", "type")
SOURCE_COLUMNS = ("source_system", "source system", "source")
RECORD_COLUMNS = ("source_record_id", "source record id", "record_id", "record id", "sys_id", "sys id")
SCENARIO_COLUMNS = ("scenario_id", "scenario id", "scenario")
RELATION_COLUMNS = ("ontology_property", "ontology property", "relationship", "relationship_type", "relationship type", "relation", "type")
START_COLUMNS = ("start_id", "start id", "source_id", "source id", ":start_id", "from_id", "from id")
END_COLUMNS = ("end_id", "end id", "target_id", "target id", ":end_id", "to_id", "to id")

SHEET_CLASS_ALIASES = {
    "approved_crs": "ChangeRecord",
    "approved_cr": "ChangeRecord",
    "crs": "ChangeRecord",
    "cr": "ChangeRecord",
    "change_requests": "ChangeRecord",
    "change_request": "ChangeRecord",
    "changes": "ChangeRecord",
    "incidents": "Incident",
    "incident": "Incident",
    "alarms": "Alarm",
    "alarm": "Alarm",
    "kpis": "KPIObservation",
    "kpi": "KPIObservation",
    "kpi_observations": "KPIObservation",
    "logs": "LogEvent",
    "log_events": "LogEvent",
    "runbooks": "Runbook",
    "runbook": "Runbook",
    "complaints": "Complaint",
    "complaint": "Complaint",
    "tickets": "TroubleTicket",
    "trouble_tickets": "TroubleTicket",
    "services": "Service",
    "service": "Service",
    "products": "Product",
    "product": "Product",
    "subscribers": "Subscriber",
    "subscriber": "Subscriber",
    "customers": "Customer",
    "customer": "Customer",
    "sites": "CellSite",
    "cell_sites": "CellSite",
    "cells": "Cell",
    "cell": "Cell",
    "routers": "Router",
    "router": "Router",
    "fiber_spans": "FiberSpan",
    "fibre_spans": "FiberSpan",
}

COLUMN_CLASS_HINTS = (
    (("change_request_id", "change request id", "cr_number", "cr number", "planned_start", "approval"), "ChangeRecord"),
    (("incident_id", "incident id", "incident_number", "incident number", "priority", "impact"), "Incident"),
    (("alarm_id", "alarm id", "probable_cause", "probable cause", "severity"), "Alarm"),
    (("kpi_name", "kpi name", "kpi_value", "kpi value", "threshold"), "KPIObservation"),
    (("runbook_id", "runbook id", "kb_number", "kb number"), "Runbook"),
    (("complaint_id", "complaint id", "ticket_id", "ticket id"), "Complaint"),
    (("service_id", "service id", "service_name", "service name"), "Service"),
)


@dataclass
class Ontology:
    classes: set[str]
    class_lookup: dict[str, str]
    properties: dict[str, dict[str, Any]]
    property_lookup: dict[str, str]


@dataclass
class IngestionBuild:
    nodes: dict[str, dict[str, Any]] = field(default_factory=dict)
    relationships: dict[tuple[str, str, str], dict[str, Any]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    files: list[IngestionFileSummary] = field(default_factory=list)


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"none", "nan"} else text


def _rel_type(property_name: str) -> str:
    parts = re.findall(r"[A-Z]?[a-z0-9]+|[A-Z]+(?=[A-Z]|$)", property_name)
    return "_".join(part.upper() for part in parts if part) or _key(property_name).upper()


def _find_column(headers: dict[str, str], candidates: tuple[str, ...]) -> str | None:
    for candidate in candidates:
        found = headers.get(_key(candidate))
        if found:
            return found
    return None


def _split_refs(value: Any) -> list[str]:
    text = _clean(value)
    if not text:
        return []
    tokens = text.splitlines() if "\n" in text else re.split(r"\s*[;,|]\s*", text)
    return [token.strip() for token in tokens if token.strip()]


class ExcelIngestionService:
    def __init__(self, project_root: Path, schema_path: Path):
        self.project_root = project_root.resolve()
        self.schema_path = schema_path.resolve()

    def ingest(self, folder_name: str, uploads: list[UploadFile]) -> IngestionResponse:
        output_dir = self._resolve_output_dir(folder_name)
        ontology = self._load_ontology()
        build = IngestionBuild()

        for upload in uploads:
            self._ingest_workbook(upload, ontology, build)

        for relationship in list(build.relationships.values()):
            self._ensure_placeholder_node(relationship[":START_ID"], self._fallback_class(ontology), build)
            self._ensure_placeholder_node(relationship[":END_ID"], self._range_class(ontology, relationship["ontology_property"]), build)

        nodes_path = output_dir / "nodes.csv"
        relationships_path = output_dir / "relationships.csv"
        self._write_csv(nodes_path, NODE_HEADERS, build.nodes.values())
        self._write_csv(relationships_path, REL_HEADERS, build.relationships.values())

        return IngestionResponse(
            folder=str(output_dir),
            nodes_path=str(nodes_path),
            relationships_path=str(relationships_path),
            node_count=len(build.nodes),
            relationship_count=len(build.relationships),
            files=build.files,
            warnings=build.warnings[:100],
        )

    def _resolve_output_dir(self, folder_name: str) -> Path:
        if not folder_name or not folder_name.strip():
            raise ValueError("Folder name is required.")
        raw_path = Path(folder_name.strip())
        output_dir = raw_path if raw_path.is_absolute() else self.project_root / raw_path
        output_dir = output_dir.resolve()
        if self.project_root not in (output_dir, *output_dir.parents):
            raise ValueError("Folder must be inside the project workspace.")
        if not output_dir.exists() or not output_dir.is_dir():
            raise ValueError("Folder does not exist.")
        if not self.schema_path.exists():
            raise ValueError("schema.json was not found at the project root.")
        return output_dir

    def _load_ontology(self) -> Ontology:
        schema = json.loads(self.schema_path.read_text(encoding="utf-8"))
        classes = {item["name"] for item in schema.get("classes", []) if item.get("name")}
        properties = {item["name"]: item for item in schema.get("object_properties", []) if item.get("name")}
        return Ontology(
            classes=classes,
            class_lookup={_key(name): name for name in classes},
            properties=properties,
            property_lookup={_key(name): name for name in properties},
        )

    def _ingest_workbook(self, upload: UploadFile, ontology: Ontology, build: IngestionBuild) -> None:
        filename = upload.filename or "uploaded.xlsx"
        if not filename.lower().endswith((".xlsx", ".xlsm")):
            build.warnings.append(f"{filename}: skipped because only .xlsx/.xlsm files are supported.")
            return
        workbook = load_workbook(io.BytesIO(upload.file.read()), data_only=True, read_only=True)
        file_rows = 0
        try:
            for sheet in workbook.worksheets:
                rows = self._sheet_rows(sheet)
                file_rows += len(rows)
                if not rows:
                    continue
                headers = {_key(header): header for header in rows[0].keys()}
                if self._is_relationship_sheet(headers):
                    self._ingest_relationship_rows(filename, sheet.title, rows, headers, ontology, build)
                else:
                    self._ingest_node_rows(filename, sheet.title, rows, headers, ontology, build)
            sheet_count = len(workbook.sheetnames)
        finally:
            workbook.close()
        build.files.append(IngestionFileSummary(filename=filename, sheets=sheet_count, rows=file_rows))

    def _sheet_rows(self, sheet: Any) -> list[dict[str, Any]]:
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(iterator, [])]
        output: list[dict[str, Any]] = []
        for values in iterator:
            row = {
                headers[index]: value
                for index, value in enumerate(values[: len(headers)])
                if index < len(headers) and headers[index]
            }
            if any(_clean(value) for value in row.values()):
                output.append(row)
        return output

    def _is_relationship_sheet(self, headers: dict[str, str]) -> bool:
        return bool(_find_column(headers, START_COLUMNS) and _find_column(headers, END_COLUMNS))

    def _ingest_node_rows(self, filename: str, sheet_name: str, rows: list[dict[str, Any]], headers: dict[str, str], ontology: Ontology, build: IngestionBuild) -> None:
        id_col = _find_column(headers, ID_COLUMNS)
        name_col = _find_column(headers, NAME_COLUMNS)
        class_col = _find_column(headers, CLASS_COLUMNS)
        source_col = _find_column(headers, SOURCE_COLUMNS)
        record_col = _find_column(headers, RECORD_COLUMNS)
        scenario_col = _find_column(headers, SCENARIO_COLUMNS)
        sheet_class = self._sheet_class(sheet_name, headers, ontology)

        for row_number, row in enumerate(rows, start=2):
            ontology_class = self._ontology_class(row.get(class_col), sheet_class, ontology)
            if not ontology_class:
                build.warnings.append(f"{filename}/{sheet_name} row {row_number}: skipped because ontology class is missing or not in schema.json.")
                continue
            canonical_id = _clean(row.get(id_col)) or _clean(row.get(name_col))
            if not canonical_id:
                build.warnings.append(f"{filename}/{sheet_name} row {row_number}: skipped because canonical_id/id is missing.")
                continue

            source_system = _clean(row.get(source_col)) or filename
            source_record_id = _clean(row.get(record_col)) or canonical_id
            scenario_id = _clean(row.get(scenario_col))
            properties = {key: value for key, value in row.items() if _clean(value)}
            build.nodes[canonical_id] = {
                "canonical_id:ID": canonical_id,
                "name": _clean(row.get(name_col)) or canonical_id,
                "ontology_class": ontology_class,
                ":LABEL": f"Entity;{ontology_class}",
                "source_system": source_system,
                "source_record_id": source_record_id,
                "source_type": sheet_name,
                "scenario_id": scenario_id,
                "properties_json": json.dumps(properties, ensure_ascii=False, default=str),
            }
            self._ingest_inline_relationships(canonical_id, ontology_class, row, headers, source_system, source_record_id, scenario_id, ontology, build)

    def _ingest_relationship_rows(self, filename: str, sheet_name: str, rows: list[dict[str, Any]], headers: dict[str, str], ontology: Ontology, build: IngestionBuild) -> None:
        start_col = _find_column(headers, START_COLUMNS)
        end_col = _find_column(headers, END_COLUMNS)
        rel_col = _find_column(headers, RELATION_COLUMNS)
        source_col = _find_column(headers, SOURCE_COLUMNS)
        record_col = _find_column(headers, RECORD_COLUMNS)
        scenario_col = _find_column(headers, SCENARIO_COLUMNS)

        for row_number, row in enumerate(rows, start=2):
            start_id = _clean(row.get(start_col))
            end_id = _clean(row.get(end_col))
            property_name = self._ontology_property(row.get(rel_col), ontology)
            if not start_id or not end_id or not property_name:
                build.warnings.append(f"{filename}/{sheet_name} row {row_number}: skipped relationship because start, end, or ontology_property is missing.")
                continue
            self._add_relationship(start_id, end_id, property_name, _clean(row.get(source_col)) or filename, _clean(row.get(record_col)) or f"{start_id}->{end_id}", _clean(row.get(scenario_col)), row, build)

    def _ingest_inline_relationships(self, canonical_id: str, ontology_class: str, row: dict[str, Any], headers: dict[str, str], source_system: str, source_record_id: str, scenario_id: str, ontology: Ontology, build: IngestionBuild) -> None:
        ignored = {_key(column) for column in (*ID_COLUMNS, *NAME_COLUMNS, *CLASS_COLUMNS, *SOURCE_COLUMNS, *RECORD_COLUMNS, *SCENARIO_COLUMNS)}
        for normalized, column in headers.items():
            if normalized in ignored:
                continue
            property_name = ontology.property_lookup.get(normalized) or ontology.property_lookup.get(re.sub(r"_ids?$", "", normalized))
            if not property_name:
                continue
            metadata = ontology.properties[property_name]
            domain_classes = set(metadata.get("domain_classes") or [])
            if domain_classes and ontology_class not in domain_classes:
                continue
            for target_id in _split_refs(row.get(column)):
                self._add_relationship(canonical_id, target_id, property_name, source_system, source_record_id, scenario_id, {column: row.get(column)}, build)

    def _add_relationship(self, start_id: str, end_id: str, property_name: str, source_system: str, source_record_id: str, scenario_id: str, properties: dict[str, Any], build: IngestionBuild) -> None:
        build.relationships[(start_id, end_id, property_name)] = {
            ":START_ID": start_id,
            ":END_ID": end_id,
            ":TYPE": _rel_type(property_name),
            "ontology_property": property_name,
            "source_system": source_system,
            "source_record_id": source_record_id,
            "scenario_id": scenario_id,
            "properties_json": json.dumps({key: value for key, value in properties.items() if _clean(value)}, ensure_ascii=False, default=str),
        }

    def _ensure_placeholder_node(self, canonical_id: str, ontology_class: str, build: IngestionBuild) -> None:
        if canonical_id in build.nodes:
            return
        build.nodes[canonical_id] = {
            "canonical_id:ID": canonical_id,
            "name": canonical_id,
            "ontology_class": ontology_class,
            ":LABEL": f"Entity;{ontology_class}",
            "source_system": "Ingestion reference",
            "source_record_id": canonical_id,
            "source_type": "relationship_endpoint",
            "scenario_id": "",
            "properties_json": json.dumps({"placeholder": True}, ensure_ascii=False),
        }

    def _ontology_class(self, value: Any, fallback: str | None, ontology: Ontology) -> str | None:
        clean = _clean(value)
        return ontology.class_lookup.get(_key(clean)) if clean else fallback

    def _sheet_class(self, sheet_name: str, headers: dict[str, str], ontology: Ontology) -> str | None:
        normalized_sheet = _key(re.sub(r"\b\d{1,2}(st|nd|rd|th)?\b", "", sheet_name, flags=re.IGNORECASE))
        direct = ontology.class_lookup.get(normalized_sheet)
        if direct:
            return direct
        alias = SHEET_CLASS_ALIASES.get(normalized_sheet)
        if alias in ontology.classes:
            return alias
        header_keys = set(headers.keys())
        for candidates, ontology_class in COLUMN_CLASS_HINTS:
            if ontology_class in ontology.classes and any(_key(candidate) in header_keys for candidate in candidates):
                return ontology_class
        return None

    def _ontology_property(self, value: Any, ontology: Ontology) -> str | None:
        clean = _clean(value)
        return ontology.property_lookup.get(_key(clean)) if clean else None

    def _range_class(self, ontology: Ontology, property_name: str) -> str:
        ranges = ontology.properties.get(property_name, {}).get("range_classes") or []
        if ranges:
            return ranges[0]
        return self._fallback_class(ontology)

    def _fallback_class(self, ontology: Ontology) -> str:
        return sorted(ontology.classes)[0] if ontology.classes else "Thing"

    def _write_csv(self, path: Path, headers: list[str], rows: Any) -> None:
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow(row)
